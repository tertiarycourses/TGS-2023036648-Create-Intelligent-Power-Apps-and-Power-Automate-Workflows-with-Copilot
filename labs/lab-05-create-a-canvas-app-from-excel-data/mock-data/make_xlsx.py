#!/usr/bin/env python3
"""Convert CoffeeMachineData.csv into CoffeeMachineData.xlsx with a FORMATTED
Excel table named CoffeeMachines.

Power Apps binds to a NAMED Excel table, not to a worksheet or a plain range.
A workbook whose data is only a range will show up in the Excel Online
(Business) connector with no tables to choose, so Lab 5 cannot proceed.

Usage:
    pip install openpyxl
    python3 make_xlsx.py
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(HERE, "CoffeeMachineData.csv")
XLSX_PATH = os.path.join(HERE, "CoffeeMachineData.xlsx")
TABLE_NAME = "CoffeeMachines"
SHEET_NAME = "CoffeeMachines"


def main() -> None:
    with open(CSV_PATH, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))

    header, data = rows[0], rows[1:]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws.append(header)
    for row in data:
        machine_id, model, location, status, last_serviced, daily_cups = row
        ws.append([machine_id, model, location, status, last_serviced, int(daily_cups)])

    last_col = get_column_letter(len(header))
    last_row = len(data) + 1
    ref = "A1:{c}{r}".format(c=last_col, r=last_row)

    table = Table(displayName=TABLE_NAME, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    widths = [12, 26, 44, 16, 16, 12]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = "0"
            cell.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "A2"

    wb.save(XLSX_PATH)
    print("Wrote {p} with table '{t}' over {r}".format(p=XLSX_PATH, t=TABLE_NAME, r=ref))


if __name__ == "__main__":
    main()
